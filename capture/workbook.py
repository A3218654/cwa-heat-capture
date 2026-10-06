"""把 CSV 紀錄整理成試算表。

- 每天資料夾：紀錄_YYYY-MM-DD.xlsx，分頁：
  縣市溫度極值／高溫資訊／熱傷害／臺北測站／體感溫度／執行狀況
- Drive 最外層：體感溫度總表.xlsx，每天一個分頁（最新的在最前面）

CSV 仍是程式追加資料用的原始紀錄；試算表每次執行後依 CSV 重新產生。
也可單獨執行補建：REBUILD_DAYS=2026-10-02,2026-10-03 python -m capture.workbook
"""
import csv
import os
import sys

from openpyxl import Workbook, load_workbook
import re

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import common, config

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# 分頁名稱 → CSV「網站」欄的關鍵字
SHEETS = [
    ("高溫資訊", "高溫資訊"),
    ("熱傷害", "熱傷害"),
    ("臺北測站", "臺北測站"),
]
DROP_COLS = {"網站"}


def _read_csv(path: str) -> tuple[list[str], list[dict]]:
    if not os.path.exists(path):
        return [], []
    with open(path, encoding="utf-8-sig") as f:
        r = csv.DictReader(f)
        return list(r.fieldnames or []), list(r)


def _write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append([_num(v) for v in row])
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] or [4])
        ws.column_dimensions[get_column_letter(i)].width = min(max(6, width * 1.6 + 2), 60)


def _num(v):
    """數字字串轉成數字，試算表才能排序計算。"""
    if isinstance(v, str):
        try:
            return int(v) if v.strip().lstrip("-").isdigit() else float(v)
        except ValueError:
            return v
    return v


# ------------------------------------------------------------------ 縣市溫度極值（使用者指定格式）
PEACH = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="D9A08B")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _one_line(v: str) -> str:
    return "、".join(x.strip() for x in str(v).splitlines() if x.strip())


def temptop_record(day: str, rows: list[dict]) -> dict | None:
    """取當天的定案值：優先用隔天截的「昨日／前日」，沒有才用當天最後一次「今日」。"""
    cand = [r for r in rows if "縣市溫度極值" in r.get("網站", "") and r.get("狀態", "").startswith("已截圖")
            and f"資料日期{day}" in r.get("補充說明", "")]
    if not cand:
        return None
    final = [r for r in cand if "(今日)" not in r.get("補充說明", "")]
    r = (final or cand)[-1]
    note = r.get("補充說明", "")
    tm = re.search(r"觀測時間([^；]+)", note)
    st = re.search(r"測站(.+?)\(", note, re.S)
    return {
        "月份": int(day[5:7]), "日期": int(day[8:10]), "最高溫": _num(r.get("數值", "")),
        "時間": _one_line(tm.group(1)) if tm else "",
        "測站": _one_line(st.group(1)) if st else "",
        "截圖": r.get("截圖檔名", ""),
    }


def write_temptop_sheet(ws, records: list[dict]) -> None:
    headers = ["月份", "日期", "最高溫", "時間", "測站"]
    ws.append(headers)
    for rec in records:
        ws.append([rec[h] for h in headers])
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=len(headers)):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.font = Font(size=12)
            c.border = BOX
            if c.row == 1:
                c.fill = PEACH
            if c.column == 3 and c.row > 1:
                c.number_format = "0.0"
    for col, w in zip("ABCDE", (10, 10, 12, 14, 18)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"


def _town_rows(day: str) -> tuple[list[str], list[list]]:
    fields, rows = _read_csv(os.path.join(common.day_dir(day), f"體感溫度_{day}.csv"))
    return fields, [[r.get(f, "") for f in fields] for r in rows]


def build_daily(day: str) -> str | None:
    fields, rows = _read_csv(common.csv_path(day))
    t_fields, t_rows = _town_rows(day)
    if not rows and not t_rows:
        return None
    wb = Workbook()
    wb.remove(wb.active)
    rec = temptop_record(day, rows)
    write_temptop_sheet(wb.create_sheet("縣市溫度極值"), [rec] if rec else [])
    cols = [f for f in fields if f not in DROP_COLS]
    for title, key in SHEETS:
        part = [r for r in rows if key in r.get("網站", "") and r.get("狀態") not in ("失敗",)]
        _write_sheet(wb.create_sheet(title), cols, [[r.get(c, "") for c in cols] for r in part])
    ws = wb.create_sheet("體感溫度")
    if t_rows:
        _write_sheet(ws, t_fields, t_rows)
    else:
        ws.append(["今天的體感溫度尚未記錄（每天 19 時後產生）"])
    bad = [r for r in rows if r.get("狀態") in ("失敗", "缺漏") or "截圖失敗" in r.get("狀態", "")]
    _write_sheet(wb.create_sheet("執行狀況"), fields, [[r.get(c, "") for c in fields] for r in bad])
    path = os.path.join(common.day_dir(day), f"紀錄_{day}.xlsx")
    wb.save(path)
    return path


def update_master(day: str) -> bool:
    t_fields, t_rows = _town_rows(day)
    if not t_rows:
        return False
    path = os.path.join(config.OUT_DIR, config.MASTER_TOWN_XLSX)
    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = Workbook()
        wb.remove(wb.active)
    if day in wb.sheetnames:
        del wb[day]
    _write_sheet(wb.create_sheet(day), t_fields, t_rows)
    wb._sheets.sort(key=lambda s: s.title, reverse=True)  # 最新的一天放最前面
    wb.active = 0
    wb.save(path)
    return True


def refresh(days: list[str]) -> list[str]:
    done = []
    for day in sorted(set(days)):
        try:
            if build_daily(day):
                done.append(f"紀錄_{day}.xlsx")
            if update_master(day):
                done.append(f"{config.MASTER_TOWN_XLSX}［{day}］")
        except Exception as e:  # noqa: BLE001  試算表失敗不影響截圖與 CSV
            print(f"[試算表] {day} 產生失敗：{e}")
    print("[試算表] 已更新：", "、".join(done) or "無")
    return done


def main() -> int:
    days = [d.strip() for d in os.environ.get("REBUILD_DAYS", "").split(",") if d.strip()]
    if not days:
        print("請設定 REBUILD_DAYS")
        return 1
    common.pull_existing(days)
    refresh(days)
    return 0 if common.push_all() else 1


if __name__ == "__main__":
    sys.exit(main())
